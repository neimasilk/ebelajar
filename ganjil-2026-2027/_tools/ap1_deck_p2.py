# -*- coding: utf-8 -*-
"""Deck Pertemuan 2 - Application Project I, kelas A (luring, Kamis 17 Sep 2026).

SUMBER (sinkron dengan Drive AP1 "pertemuan 2", folder 1yvgECF7RK8l6Hr4TG0oIG-RFdIkocAou):
  - "Modul Pertemuan 2.pdf" (Modul Mandiri): ergonomi & aksesibilitas + ideasi proyek.
    Peta: pemanasan -> audit wireframe P1 -> perbaikan terarah -> ideasi -> refleksi.
    Praktik A (audit, checklist 6 butir, tabel >=5 temuan), B (revisi 3-5 perbaikan),
    C (ide proyek, format 7 butir). Penugasan 2a/2b, rubrik 20/30/25/15/10 lulus >=70,
    refleksi 3 soal, mini kuis 3 soal (kunci 1-b 2-b 3-c), rujukan Olsen/Ries/Knapp.
  - Video "Bedah Ergonomi & Aksesibilitas UI" dan "Rahasia Desain Aplikasi".
  Pengayaan dari modul_p2.html lokal (Fitts's Law, target sentuh 44 px, zona ibu jari,
  WCAG 4.5:1, jangan andalkan warna, grayscale / zoom 200% / urutan fokus).

Keputusan user 17 Sep: Tugas 2 (Brainstorming, 6%) DIMAJUKAN dari P3 ke P2, isinya
mengikuti modul (2a audit + revisi wireframe individu, 2b ide proyek kelompok).
Slide #2-4 = umpan balik Pertemuan 1 (tugas, forum, perkenalan, kelompok) - angka di
konstanta ASOF/T1_KUMPUL/FORUM_P1_POST/NILAI_TYO/NILAI_DEVAN atau argumen CLI.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import *
from deckkit import _tb, _emit

OUT = r"D:\documents\ebelajar\ganjil-2026-2027\application-project-1"
B = "https://ebelajar.stiki.ac.id"

# -- cmid AP1 A (course 7278; sweep 17 Sep) --
COURSE     = 7278
KONTRAK    = 496744
MATERI2    = 491628   # page "Ergonomi dan Aksesibilitas dalam Desain UI + Pencarian Ide Proyek"
TUGAS1     = 491624   # assign "TUGAS 1" (tenggat 20 Sep 23.55)
TUGAS2     = 491632   # assign "Tugas 2" (dipindah ke Section 2)
FORUM_P1   = 491623
PERKENALAN = 491618
UMUM       = 491616
EXIT_P2    = 502324   # forum "Exit-ticket Pertemuan 2" (dibuat 17 Sep, Section 2)
SECTION2   = "%s/course/view.php?id=%d#section-2" % (B, COURSE)

# -- Drive AP1 --
FOLDER_P2 = "https://drive.google.com/drive/folders/1yvgECF7RK8l6Hr4TG0oIG-RFdIkocAou"
MODUL_P2  = "https://drive.google.com/file/d/1N1ZWqMpPIyUTYAOYnecxqnIhz8VCz2Un/view"
VIDEO_1   = "https://drive.google.com/file/d/1UHe79xxVM1WRlpUSByVUziT7NcGm09oW/view"
VIDEO_2   = "https://drive.google.com/file/d/1rXAx7UDSIC38bylwRptd4GE4od9Dimtq/view"
SHEET_KEL = "https://docs.google.com/spreadsheets/d/12mbDubq6xxfcrlFEQir4aG4h-ICjM0g2B9EU1IuzJUU/edit"

TENGGAT = "Minggu, 27 September 2026, pukul 23.55"

# -- Umpan balik Pertemuan 1 (ubah angka lalu jalankan ulang; bisa juga lewat argumen CLI,
#    mis. `python ap1_deck_p2.py --tyo 85 --devan 82 --t1 4 --forum-p1 0`) --
ASOF          = "Kamis 17 Sep, 10.40"   # waktu sweep eBelajar + sheet
PESERTA       = 21
T1_KUMPUL     = 0      # Tugas 1 yang sudah dikumpulkan
FORUM_P1_POST = 0      # posting mahasiswa di Diskusi Pertemuan 1
NILAI_TYO     = 80     # nilai rating forum Perkenalan (0-100)
NILAI_DEVAN   = 80


def build():
    prs = new_deck()

    title_slide(prs, "PERTEMUAN 2 · APPLICATION PROJECT I",
                "Nyaman Dipakai,\nTerbuka untuk Semua",
                "Ergonomi & aksesibilitas UI — lalu mencari ide proyek kelompok",
                "Mukhlis Amien, M.Kom.  ·  IF24KB33  ·  Kelas A (Reguler) · Kamis 13.50–15.30 · Ruang A.2.1")

    # ══ UMPAN BALIK PERTEMUAN 1 ══
    if T1_KUMPUL == 0:
        t1 = "%d dari %d mengumpulkan · tenggat Minggu 20 Sep 23.55 — belum ada yang bisa dikoreksi" % (T1_KUMPUL, PESERTA)
    else:
        t1 = "%d dari %d mengumpulkan · tenggat Minggu 20 Sep 23.55 — koreksi menyusul setelah tenggat" % (T1_KUMPUL, PESERTA)
    if FORUM_P1_POST == 0:
        fp1 = "Belum ada posting mahasiswa — forum tetap terbuka dan dinilai"
    else:
        fp1 = "%d posting — dibalas dan dinilai di eBelajar" % FORUM_P1_POST
    table_slide(prs, "Umpan balik Pertemuan 1",
                ["Aktivitas", "Kondisi per %s" % ASOF],
                [["*(a) Tugas 1 — Wireframe", t1],
                 ["*(b) Forum Diskusi P1", fp1],
                 ["*(c) Forum Perkenalan", "Tyo dan Devan sudah mengisi → dibalas dan dinilai di eBelajar "
                                           "(Tyo %s · Devan %s)" % (NILAI_TYO, NILAI_DEVAN)],
                 ["*(d) Kelompok proyek", "4 kelompok sah (1 belum bernama) · 9 orang belum berkelompok → dikunci hari ini"]],
                col_w=[3.6, 8.4], fs=17,
                sub="Koreksi dan komentar atas pekerjaan minggu lalu — kondisi nyata di eBelajar.",
                note="Nada: bukan menegur, tapi memperlihatkan angka. Tugas 1 masih ada waktu sampai "
                     "Minggu. Forum Perkenalan dan Diskusi P1 masih terbuka dan tetap dinilai "
                     "(komponen Partisipatif 10%). Rincian catatan di slide berikut.")

    content_slide(prs, "Catatan umpan balik",
                  ["**Tugas 1** — sebelum mengunggah, cek dua hal yang paling sering terlupa:",
                   (1, "**Anotasi**: panah + satu kalimat *kenapa* elemen itu diletakkan di sana"),
                   (1, "**Paragraf alasan**: hierarki visual, konsistensi, keterbacaan — tulis alasannya, bukan deskripsi gambarnya"),
                   (1, "Bobot Tugas 1 = **3%** sesuai Kontrak Kuliah (slide P1 sempat menulis 5% — itu angka RPS per blok minggu)"),
                   "**Perkenalan** — yang masuk baru memuat **identitas administratif** (nama, asal, status, prodi, NRP)",
                   (1, "Perkenalan yang berguna untuk membentuk kelompok menambahkan: **minat** (desain, front-end, back-end), "
                       "**pengalaman** (proyek atau alat yang pernah dipakai), dan **harapan** dari proyek semester ini"),
                   (1, "*Contoh: \"Lebih nyaman di front-end, pernah membuat landing page dengan HTML/CSS, ingin proyek yang "
                       "mengurangi antrean di kantin kampus.\"*"),
                   "__",
                   "Perkenalan yang sudah masuk boleh **disunting untuk dilengkapi** — dan yang belum, silakan mengisi dengan format ini."],
                  body_size=18,
                  note="Contoh perkenalan di atas hanya ilustrasi format, bukan kutipan mahasiswa.")

    ks = table_slide(prs, "Kelompok: status per 17 Sep",
                ["Kelompok", "Anggota tercatat di sheet KELOMPOK AP 1", "Status"],
                [["*RAWON", "Andrean (ketua), Jovan, Jeremiah, Devan", "Sah · 4 — penuh"],
                 ["*JOSJIS", "Bobby (ketua), Tyo, Geraldy", "Sah · 3 — sisa 1 kursi"],
                 ["*QWERTY", "Daffa (ketua — NRP belum ditulis), Marchel, Jibril", "Sah · 3 — sisa 1 kursi"],
                 ["*(belum bernama)", "Rizkia, Kirana", "Sah · 2 — butuh nama; sisa 2 kursi"],
                 ["*Sendiri", "Notopianus Muyapa · Tekiu Newegalen", "Belum berkelompok"],
                 ["*Belum tercatat", "Firmanda Sulistyono, Mochammad Feyza Maulana, Shidqi Andriyanto Darman, "
                                    "Yudha Abi Krisnanda, Felix Khancitra Subekti, Setya Iqbal Putra Candra, "
                                    "Ferdi Alexander Kiwak", "Belum berkelompok (7 orang)"]],
                col_w=[2.3, 6.6, 3.1], fs=14,
                sub="Aturan: minimal 2, maksimal 4 orang per kelompok, tetap sampai akhir semester.",
                note="Aritmetika: 4 kursi kosong di kelompok lama (JOSJIS 1, QWERTY 1, Rizkia-Kirana 2). "
                     "9 orang belum berkelompok -> kalau 4 kursi terisi, 5 sisanya membentuk kelompok baru "
                     "(mis. 3 + 2); atau 9 orang menjadi 3 kelompok baru. Kunci di 10 menit pertama studio.")
    tf = _tb(ks, 0.85, 5.75, 11.65, 1.4)
    for i, line in enumerate([
            "**Aritmetika:** 4 kursi kosong di kelompok lama → minimal **5 orang membentuk kelompok baru** (mis. 3 + 2) — atau 9 orang menjadi 3 kelompok baru.",
            "**Aksi:** dikunci di kelas hari ini (10 menit pertama studio) → sheet diperbarui **sebelum kelas bubar** → unggah Tugas 1b. "
            "Kolom **Topik 1–3** semua kelompok masih kosong → isi sebelum Tugas 2b."]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_before = Pt(0 if i == 0 else 6)
        _emit(para, line, 15, INK, NAVY)

    content_slide(prs, "Alur 100 menit hari ini",
                  ["**Hook** — tombol kecil itu biaya, bukan selera · 5'",
                   "**Mini-lecture** — lima prinsip ergonomi & aksesibilitas + tiga cek cepat · 25'",
                   "**Studio** — kunci kelompok → audit wireframe P1 → revisi → ideasi proyek · 55'",
                   "**Debrief** — dua kelompok memaparkan satu ide dalam 60 detik · 10'",
                   "**Exit-ticket** — tiga pertanyaan refleksi di eBelajar · 5'",
                   "__",
                   "Bawa wireframe Tugas 1 Anda — kertas atau Figma. Itu bahan kerja studio hari ini."],
                  sub="Urutan ini mengikuti peta Modul Mandiri Pertemuan 2 di Google Drive.",
                  body_size=20)

    # ══ HOOK ══
    statement_slide(prs, "Tombol yang terlalu kecil\nbukan soal selera.",
                    "Itu biaya yang dibayar pengguna — setiap kali ia meleset, mengulang, atau menyerah. "
                    "Hari ini kita belajar melihat biaya itu, lalu memperbaikinya di wireframe Anda sendiri.",
                    note="Hook 5 menit. Tanya kelas: kapan terakhir salah pencet di ponsel? Tombol apa, "
                         "di aplikasi apa, dan kenapa? Kumpulkan 2–3 jawaban, lalu sambungkan: itu "
                         "masalah ergonomi, bukan kecerobohan pengguna.")

    content_slide(prs, "Capaian pertemuan ini (Sub-CPMK0922)",
                  ["Menjelaskan prinsip **ergonomi digital** dan **aksesibilitas** pada UI berbasis web",
                   "Menerapkan prinsip tersebut untuk **memperbaiki wireframe** hasil Pertemuan 1",
                   "Menghasilkan **1–2 ide proyek awal**: definisi masalah, target pengguna, hipotesis nilai",
                   "Menyusun **rencana langkah berikutnya** untuk validasi ide di Pertemuan 3",
                   "__",
                   "**Selesai berarti:** audit + wireframe revisi terunggah, dokumen ide proyek terunggah, rubrik ≥ 70/100."],
                  sub="Empat tujuan ini diambil langsung dari Modul Mandiri P2 — dan menjadi dasar Tugas 2.",
                  body_size=20)

    content_slide(prs, "Bahan belajar di Drive — Pertemuan 2",
                  ["**Modul Mandiri Pertemuan 2 (PDF)** — materi, tiga praktik terarah, penugasan, rubrik, kuis",
                   "Video **Bedah Ergonomi & Aksesibilitas UI** — dari ide proyek hingga perbaikan",
                   "Video **Rahasia Desain Aplikasi**",
                   "Estimasi belajar mandiri: **120–150 menit**",
                   "Rujukan: Olsen, *The Lean Product Playbook* (2015) · Ries, *The Lean Startup* (2011) · Knapp dkk., *Sprint* (2016)"],
                  sub="Tautan ada di slide terakhir. eBelajar hanya menautkan — satu salinan di Drive.",
                  body_size=20)

    # ══ BAGIAN 1 — MATERI ══
    section_slide(prs, "BAGIAN 1 · 25 MENIT", "Ergonomi & Aksesibilitas",
                  "Ergonomi: nyaman, efisien, aman dipakai · Aksesibilitas: bisa dipakai sebanyak mungkin orang")

    content_slide(prs, "Prinsip 1 — Target sentuh dan spasi",
                  ["Elemen interaktif **mudah disentuh** dan **tidak berdempetan** berlebihan",
                   "Jarak dan ukuran tombol **konsisten** antar layar",
                   "__",
                   "**Fitts's Law** — makin besar dan makin dekat sebuah target, makin cepat ia dicapai",
                   (1, "Tombol aksi utama: besar, dan dekat tempat mata/jari pengguna berada"),
                   "Target sentuh **≥ 44 px** — tombol kecil adalah biaya bagi pengguna",
                   "**Zona ibu jari** — di ponsel, area paling nyaman ada di bagian bawah layar; aksi penting jangan di pojok atas"],
                  body_size=19,
                  note="Peragakan: pegang ponsel satu tangan, coba jangkau pojok kiri atas. Itu alasan "
                       "banyak aplikasi memindahkan navigasi utama ke bawah.")

    content_slide(prs, "Prinsip 2 — Konteks dan umpan balik",
                  ["Sistem memberi **umpan balik jelas** untuk setiap klik, proses, dan error",
                   (1, "Tombol ditekan → terlihat ditekan"),
                   (1, "Proses berjalan → ada tanda sedang memuat"),
                   (1, "Gagal → pesan yang menjelaskan apa yang terjadi"),
                   "**Navigasi dan status langkah** mudah dilacak pengguna",
                   (1, "Di mana saya sekarang? Tinggal berapa langkah lagi?"),
                   "__",
                   "Diam adalah umpan balik terburuk: pengguna menekan ulang, lalu datanya terkirim dua kali."],
                  body_size=19)

    content_slide(prs, "Prinsip 3 — Beban kognitif rendah",
                  ["**Hindari pilihan berlebihan** pada satu layar",
                   "Gunakan **pola yang dapat diprediksi** dan **istilah yang familiar**",
                   "Tugas yang sama selesaikan dengan **langkah seringan mungkin**",
                   (1, "Setiap langkah ekstra adalah kesempatan pengguna menyerah"),
                   "__",
                   "Uji sederhana: bisakah orang yang **baru pertama kali** membuka layar ini menyebutkan apa yang harus ia lakukan — dalam lima detik?"],
                  body_size=20)

    content_slide(prs, "Prinsip 4 — Responsif dan toleran terhadap error",
                  ["Tata letak **menyesuaikan berbagai ukuran layar** — ponsel, tablet, laptop",
                   "Form **mencegah kesalahan** sebelum terjadi",
                   (1, "Format jelas di awal, pilihan dibatasi bila memang terbatas"),
                   "Saat salah, beri **saran perbaikan yang spesifik**",
                   (1, "Bukan \"Input tidak valid\", melainkan \"Nomor HP diawali 08 dan berisi 10–13 angka\""),
                   "__",
                   "Relevan langsung dengan stack semester ini: komponen **React** yang Anda bangun di Pertemuan 4–5 harus mewarisi aturan ini."],
                  body_size=19)

    content_slide(prs, "Prinsip 5 — Aksesibilitas dasar",
                  ["**Kontras memadai** antara teks dan latar — acuan WCAG: **4.5:1** untuk teks biasa",
                   "**Jangan andalkan warna saja** — status merah/hijau tidak terbaca oleh pengguna buta warna; tambahkan ikon atau teks",
                   "**Teks alternatif** untuk gambar yang bermakna — pembaca layar menggantikan mata pengguna",
                   "**Fokus keyboard** dan **urutan navigasi yang logis** — banyak orang tidak memakai mouse",
                   "**Ukuran huruf fleksibel** — rancangan yang rusak saat huruf diperbesar adalah rancangan yang rapuh"],
                  sub="Aksesibilitas bukan fitur tambahan. Ia kualitas yang menentukan siapa yang bisa masuk.",
                  body_size=19)

    table_slide(prs, "Tiga cek cepat — lakukan pada wireframe Anda",
                ["Cek", "Caranya", "Yang dicari"],
                [["*Grayscale", "Lihat rancangan tanpa warna", "Apakah informasi tetap terbaca?"],
                 ["*Zoom 200%", "Perbesar tampilan dua kali lipat", "Apakah tata letak bertahan?"],
                 ["*Urutan fokus", "Telusuri elemen dengan tombol Tab", "Apakah urutannya masuk akal?"]],
                col_w=[2.5, 4.5, 5], fs=18,
                sub="Tiga cek ini murah, cepat, dan menemukan sebagian besar masalah sebelum ada satu baris kode pun.",
                note="Untuk wireframe kertas: cek urutan fokus dengan menomori elemen interaktif "
                     "sesuai urutan yang Anda harapkan, lalu minta teman menelusuri.")

    content_slide(prs, "Mini kuis — jawab dalam hati, 1 menit",
                  ["**1.** Aksesibilitas terutama memastikan bahwa…",
                   (1, "a) warna terlihat cerah · b) aplikasi dapat digunakan oleh beragam pengguna · c) ukuran file kecil · d) server cepat"),
                   "**2.** Indikator beban kognitif rendah:",
                   (1, "a) terlalu banyak opsi dalam satu layar · b) pola interaksi konsisten · c) variasi gaya komponen berlebihan · d) istilah teknis tanpa penjelasan"),
                   "**3.** Saat audit ergonomi, prioritas pertama adalah…",
                   (1, "a) memperbanyak ikon · b) menambah teks · c) memperbaiki navigasi dan aksi utama · d) membuat animasi")],
                  sub="Soal dari Modul Mandiri P2. Kunci dibahas setelah semua menjawab.",
                  body_size=19,
                  note="Kunci: 1-b, 2-b, 3-c. Soal 3 adalah jembatan ke studio: saat merevisi "
                       "wireframe, mulai dari navigasi dan aksi utama.")

    # ══ BAGIAN 2 — STUDIO ══
    section_slide(prs, "BAGIAN 2 · 55 MENIT", "Studio",
                  "Kunci kelompok → audit → revisi → ideasi proyek")

    table_slide(prs, "Studio hari ini — 55 menit",
                ["Menit", "Kegiatan", "Hasil"],
                [["*0–10", "Kunci kelompok (2–4 orang) — perbarui sheet KELOMPOK AP 1", "Nama, ketua, anggota final"],
                 ["*10–22", "Praktik A — audit ergonomi wireframe P1", "Tabel ≥ 5 temuan"],
                 ["*22–35", "Praktik B — revisi wireframe", "3–5 perbaikan + catatan"],
                 ["*35–55", "Praktik C — ideasi proyek per kelompok", "1–2 ide, format 7 butir"]],
                col_w=[1.6, 6.4, 4], fs=17,
                sub="Di modul, tiap praktik 30–45 menit. Di kelas kita mulai bersama; sisanya selesai sebagai Tugas 2.",
                note="Berkeliling per kelompok. Di Praktik A-B pertanyaan kuncinya: apa aksi utama "
                     "di wireframe ini, dan bisakah dijangkau ibu jari? Di Praktik C: siapa "
                     "penggunanya, dan apa asumsi yang paling mungkin salah?")

    table_slide(prs, "Praktik A — daftar cek audit ergonomi",
                ["#", "Pertanyaan untuk wireframe P1 Anda"],
                [["*1", "Apakah tombol utama terlihat menonjol dan mudah dijangkau?"],
                 ["*2", "Apakah jarak antar elemen interaktif cukup?"],
                 ["*3", "Apakah navigasi selalu berada di lokasi yang sama di setiap layar?"],
                 ["*4", "Apakah pesan error dan status sistem tampak dan mudah dipahami?"],
                 ["*5", "Apakah kontras teks terhadap latar memadai?"],
                 ["*6", "Apakah urutan tab/fokus mudah diprediksi?"]],
                col_w=[1, 11], fs=18,
                sub="Buka atau cetak wireframe Tugas 1. Nilai satu per satu — jangan dari ingatan.")

    table_slide(prs, "Praktik A — bentuk tabel temuan (minimal 5 baris)",
                ["Masalah", "Dampak", "Usulan perbaikan"],
                [["Tiga tombol sama besar di beranda", "Pengguna ragu mana aksi utamanya", "Satu tombol utama dibuat paling menonjol"],
                 ["Tombol keranjang di pojok kanan atas", "Sulit dijangkau ibu jari di ponsel", "Pindahkan ke bilah bawah"],
                 ["Teks abu muda di latar putih", "Sulit dibaca, terutama di luar ruangan", "Pertegas warna teks sampai kontras cukup"],
                 ["…", "…", "…"]],
                col_w=[4, 4, 4], fs=16,
                sub="Contoh baris di bawah hanya ilustrasi format — temuan Anda harus berasal dari wireframe Anda sendiri.",
                note="Tekankan kolom Dampak: temuan tanpa dampak hanyalah opini. Rubrik menilai "
                     "kelengkapan dan ketajaman temuan (20 poin).")

    content_slide(prs, "Praktik B — revisi wireframe",
                  ["**Pilih 3–5 perbaikan prioritas** dari tabel audit — tidak semuanya harus dikerjakan",
                   (1, "Mulai dari **navigasi dan aksi utama** (ingat kuis nomor 3)"),
                   "**Revisi wireframe:**",
                   (1, "perjelas hierarki"),
                   (1, "perbesar area interaktif yang perlu"),
                   (1, "perbaiki kontras"),
                   "**Tambahkan catatan** pada wireframe: apa yang berubah, dan **kenapa**",
                   "__",
                   "Output: wireframe revisi + catatan perubahan (PDF/PNG). Rubrik: 30 poin — komponen terbesar."],
                  body_size=19)

    table_slide(prs, "Praktik C — format ide proyek",
                ["#", "Bagian", "Isi"],
                [["*1", "Judul sementara", "Nama kerja + deskripsi singkat"],
                 ["*2", "Pernyataan masalah", "Siapa pengguna, konteksnya, kendala utamanya"],
                 ["*3", "Target pengguna", "Persona mini: tujuan utama, kebiasaan, perangkat utama"],
                 ["*4", "Hipotesis nilai", "Manfaat utama yang dijanjikan kepada pengguna"],
                 ["*5", "Fitur inti", "3–5 poin yang langsung mengatasi masalah inti"],
                 ["*6", "Risiko & asumsi kritis", "Apa yang paling perlu diverifikasi lebih dulu"],
                 ["*7", "Metrik awal keberhasilan", "Contoh metrik sederhana yang bisa diukur"]],
                col_w=[0.8, 3.8, 7.4], fs=16,
                sub="1–2 ide per kelompok, maksimal 1 halaman per ide. Sertakan daftar anggota dan peran singkat.")

    content_slide(prs, "Ide yang baik bukan yang paling keren",
                  ["**Masalah nyata** — dialami sendiri atau lingkungan sekitar; hindari klon aplikasi populer tanpa alasan",
                   "**Pengguna jelas** — satu kalimat: siapa dia, dan apa yang ingin ia capai",
                   "**Fitur inti sedikit** — yang paling langsung menyelesaikan masalah, bukan daftar sepanjang tangan",
                   "**Uji paling murah dulu** — wireframe dan calon pengguna, sebelum menulis kode",
                   "**Sesuai kemampuan tim dan ruang lingkup** — dibangun sebagai web dengan **React + Node.js** dalam satu semester",
                   "__",
                   "Kerangka ini mengikuti Dan Olsen, *The Lean Product Playbook* — rujukan utama Tugas 1–2."],
                  body_size=19,
                  note="Isi kolom Topik 1–3 di sheet KELOMPOK AP 1 dengan judul sementara ide-ide ini.")

    # ══ BAGIAN 3 — TUGAS 2 ══
    section_slide(prs, "BAGIAN 3", "Tugas 2",
                  "Audit Ergonomi & Ide Proyek (Brainstorming) · bobot 6%")

    content_slide(prs, "Tugas 2 — apa yang dikumpulkan",
                  ["**2a — Audit ergonomi & revisi wireframe (individu)**",
                   (1, "Tabel audit minimal 5 temuan → **AP1_P2_NIM_Nama_Audit.pdf**"),
                   (1, "Wireframe revisi + catatan perubahan → **AP1_P2_NIM_Nama_WireframeRevisi.pdf**"),
                   "**2b — Ide proyek awal (kelompok)**",
                   (1, "1–2 ide sesuai format 7 butir + daftar anggota dan peran → **AP1_P2_KelompokX_IdeProyek.pdf**"),
                   (1, "**Setiap anggota mengunggah berkas 2b yang sama.** Aturan \"berkas identik = 0\" **tidak berlaku** untuk berkas kelompok ini"),
                   "__",
                   "Semua lewat **eBelajar → Pertemuan 2 → Tugas 2**, maksimal 5 berkas. Bukan WhatsApp, bukan email."],
                  body_size=19)

    table_slide(prs, "Rubrik Tugas 2 — total 100",
                ["Komponen", "Poin"],
                [["*Audit ergonomi: kelengkapan dan ketajaman temuan", "20"],
                 ["*Kualitas perbaikan wireframe: penerapan ergonomi & aksesibilitas", "30"],
                 ["*Kejelasan ide proyek: masalah, target pengguna, hipotesis nilai", "25"],
                 ["*Kelayakan dan fokus fitur inti", "15"],
                 ["*Kerapian penyajian dan kepatuhan format", "10"]],
                col_w=[10, 2], fs=18,
                sub="Kriteria lulus tugas: ≥ 70 poin. Nilai Tugas 2 menyumbang 6% nilai akhir (komponen Observasi).",
                note="Rubrik identik dengan Modul Mandiri P2.")

    content_slide(prs, "Tenggat",
                  ["Dibuka **hari ini**, Kamis 17 September 2026",
                   "Tenggat: **%s**" % TENGGAT,
                   "Terlambat: **−10% per hari**, maksimal 3 hari; lewat itu dinilai 0 tapi **tetap wajib dikumpulkan**",
                   "__",
                   "**Catatan jadwal:** di Kontrak Kuliah, Tugas 2 semula tercantum di Pertemuan 3. Dimajukan ke Pertemuan 2 supaya sejalan dengan modul. **Bobotnya tetap 6%**, total tetap 100%.",
                   "Jangan lupa: **Tugas 1 tetap bertenggat Minggu, 20 September 2026, 23.55** — wireframe itulah bahan audit Tugas 2."],
                  body_size=19)

    content_slide(prs, "Debrief — 10 menit",
                  ["**Dua kelompok** memaparkan **satu ide** dalam **60 detik**: masalah, pengguna, fitur inti",
                   "Kelas menjawab satu pertanyaan untuk tiap ide:",
                   (1, "**Asumsi mana yang paling mungkin salah?**"),
                   "Satu orang memperlihatkan **perbaikan wireframe** terbesarnya — sebelum dan sesudah",
                   "__",
                   "Tujuannya bukan memilih ide terbaik — melainkan belajar melihat risiko sebelum membangun."],
                  body_size=20)

    content_slide(prs, "Exit-ticket — sebelum keluar ruangan",
                  ["Buka **eBelajar → Pertemuan 2 → Exit-ticket Pertemuan 2** (tautan di slide terakhir), balas diskusinya:",
                   (1, "Perbaikan apa yang paling **mengurangi beban kognitif** pada wireframe Anda?"),
                   (1, "**Asumsi kritis** apa dalam ide proyek yang harus diverifikasi lebih dulu?"),
                   (1, "Jika harus **menghapus satu fitur**, mana yang paling sedikit dampaknya pada nilai untuk pengguna?"),
                   "__",
                   "Jawaban forum masuk komponen **Partisipatif (10%)**. Forum Perkenalan dan Diskusi Pertemuan 1 juga masih terbuka."],
                  body_size=19,
                  note="Tiga soal = Refleksi Cepat Modul Mandiri P2. Kalau waktu sempit, soal 1 di "
                       "kelas, soal 2–3 dilanjutkan di rumah sebelum Minggu.")

    content_slide(prs, "Pertemuan 3 — Estetika dalam desain UI",
                  ["Komposisi dan tata letak, **warna**, **tipografi**, ritme visual dan ruang kosong",
                   "Mengubah wireframe revisi menjadi **mockup** yang lebih matang",
                   "Memadatkan ide proyek menjadi **draft proposal** ringkas",
                   "__",
                   "**Siapkan:** akun Figma aktif · wireframe revisi Tugas 2 · ide proyek kelompok yang sudah disepakati",
                   "Kendala teknis → **forum Diskusi Umum**, bukan japri",
                   "__",
                   "Sampai jumpa **Kamis depan, 24 September 2026**."],
                  body_size=20)

    links_slide(prs, "Tautan cepat — klik langsung dari slide ini",
                [("Forum Exit-ticket Pertemuan 2 — jawab sebelum keluar", "%s/mod/forum/view.php?id=%d" % (B, EXIT_P2)),
                 ("Materi P2: Ergonomi dan Aksesibilitas + Pencarian Ide Proyek", "%s/mod/page/view.php?id=%d" % (B, MATERI2)),
                 ("Tugas 2 — unggah di sini", "%s/mod/assign/view.php?id=%d" % (B, TUGAS2)),
                 ("TUGAS 1 — tenggat Minggu 20 Sep 23.55", "%s/mod/assign/view.php?id=%d" % (B, TUGAS1)),
                 ("Sheet KELOMPOK AP 1", SHEET_KEL),
                 ("Folder Drive Pertemuan 2 (modul & video)", FOLDER_P2),
                 ("Modul Mandiri Pertemuan 2 (PDF)", MODUL_P2),
                 ("Video: Bedah Ergonomi & Aksesibilitas UI", VIDEO_1),
                 ("Video: Rahasia Desain Aplikasi", VIDEO_2),
                 ("Forum Diskusi Umum — tempat bertanya", "%s/mod/forum/view.php?id=%d" % (B, UMUM))],
                sub="Salindia ini tersimpan di Google Drive; semua tautan aktif saat dibuka di PowerPoint atau Google Slides.")

    os.makedirs(OUT, exist_ok=True)
    p = os.path.join(OUT, "AP1_A_Pertemuan2_Ganjil2026.pptx")
    prs.save(p)
    return p


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--tyo"); ap.add_argument("--devan")
    ap.add_argument("--t1", type=int); ap.add_argument("--forum-p1", type=int)
    ap.add_argument("--asof")
    a = ap.parse_args()
    if a.tyo is not None: NILAI_TYO = a.tyo
    if a.devan is not None: NILAI_DEVAN = a.devan
    if a.t1 is not None: T1_KUMPUL = a.t1
    if a.forum_p1 is not None: FORUM_P1_POST = a.forum_p1
    if a.asof is not None: ASOF = a.asof
    path = build()
    from pptx import Presentation
    print("OK", path, "slides:%d" % len(Presentation(path).slides._sldIdLst))
